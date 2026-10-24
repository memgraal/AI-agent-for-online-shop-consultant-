"""Extract bounded text from supported uploads for the in-memory knowledge index."""

from io import BytesIO
from pathlib import Path
import zipfile
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from app.config import settings


class UnsupportedDocumentType(ValueError):
    """Raised when an uploaded file format is not supported for text extraction."""


class InvalidDocument(ValueError):
    """Raised when a supported file is malformed or contains no extractable text."""


class DocumentTooLarge(ValueError):
    """Raised when upload bytes or extracted text exceed configured limits."""


def extract_document_pages(filename: str, content: bytes) -> list[tuple[int | None, str]]:
    """Extract text and page numbers from PDF, DOCX, Markdown, and UTF-8 text."""
    if not content:
        raise InvalidDocument("Файл пуст.")
    if len(content) > settings.knowledge_max_upload_bytes:
        raise DocumentTooLarge("Размер файла превышает лимит загрузки.")

    suffix = Path(filename).suffix.casefold()
    if suffix in {".txt", ".md"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InvalidDocument("Текстовый файл должен быть в UTF-8.") from exc
        if "\x00" in text:
            raise InvalidDocument("Текстовый файл содержит бинарные данные.")
        pages = [(None, text)]
    elif suffix == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise InvalidDocument("Файл не похож на PDF.")
        try:
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise InvalidDocument("Зашифрованные PDF не поддерживаются.")
            if len(reader.pages) > settings.knowledge_max_pdf_pages:
                raise DocumentTooLarge("PDF содержит слишком много страниц.")
            pages = [(index, page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
        except DocumentTooLarge:
            raise
        except InvalidDocument:
            raise
        except Exception as exc:
            raise InvalidDocument("Не удалось прочитать PDF.") from exc
    elif suffix == ".docx":
        if not zipfile.is_zipfile(BytesIO(content)):
            raise InvalidDocument("Файл не похож на DOCX.")
        try:
            with ZipFile(BytesIO(content)) as archive:
                max_uncompressed_bytes = settings.knowledge_max_text_chars * 8 + 1_048_576
                if sum(entry.file_size for entry in archive.infolist()) > max_uncompressed_bytes:
                    raise DocumentTooLarge("DOCX содержит слишком много распаковываемых данных.")
            document = Document(BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            blocks.extend(
                "\t".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            pages = [(None, "\n\n".join(blocks))]
        except DocumentTooLarge:
            raise
        except Exception as exc:
            raise InvalidDocument("Не удалось прочитать DOCX.") from exc
    else:
        raise UnsupportedDocumentType("Поддерживаются только .txt, .md, .pdf и .docx.")

    extracted_chars = sum(len(text) for _, text in pages)
    if extracted_chars > settings.knowledge_max_text_chars:
        raise DocumentTooLarge("Извлечённый текст превышает допустимый размер.")
    if not any(text.strip() for _, text in pages):
        raise InvalidDocument("В файле не найден извлекаемый текст.")
    return pages
