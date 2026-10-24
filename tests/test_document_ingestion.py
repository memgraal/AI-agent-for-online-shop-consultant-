from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
import pytest

from app.services.document_ingestion import (
    DocumentTooLarge,
    InvalidDocument,
    UnsupportedDocumentType,
    extract_document_pages,
)


def test_extracts_utf8_text_and_strips_bom():
    result = extract_document_pages("manual.txt", b"\xef\xbb\xbfPower: 12 V")

    assert result == [(None, "Power: 12 V")]


def test_extracts_docx_paragraphs_and_table_content():
    document = Document()
    document.add_paragraph("Инструкция по установке")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Питание"
    table.cell(0, 1).text = "12 В"
    buffer = BytesIO()
    document.save(buffer)

    result = extract_document_pages("manual.docx", buffer.getvalue())

    assert result == [(None, "Инструкция по установке\n\nПитание\t12 В")]


def test_rejects_unsupported_file_extensions():
    with pytest.raises(UnsupportedDocumentType):
        extract_document_pages("manual.xls", b"content")


def test_rejects_invalid_pdf_signature():
    with pytest.raises(InvalidDocument, match="не похож на PDF"):
        extract_document_pages("manual.pdf", b"not a pdf")


def test_rejects_empty_or_non_utf8_text():
    with pytest.raises(InvalidDocument, match="пуст"):
        extract_document_pages("empty.txt", b"")
    with pytest.raises(InvalidDocument, match="UTF-8"):
        extract_document_pages("manual.txt", b"\xff\xfe")


def test_rejects_files_over_configured_limit(monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "knowledge_max_upload_bytes", 1024)

    with pytest.raises(DocumentTooLarge):
        extract_document_pages("manual.txt", b"a" * 1025)


def test_rejects_docx_with_excessive_uncompressed_size(monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "knowledge_max_text_chars", 1000)
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 1_100_000)

    with pytest.raises(DocumentTooLarge, match="распаковываемых данных"):
        extract_document_pages("compressed.docx", buffer.getvalue())
